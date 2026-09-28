"""校验模块：检查输出文件的完整性。"""

from pathlib import Path


def validate(output_path: str | Path, source_format: str, target_format: str) -> dict:
    """校验转换结果。

    参数:
        output_path: 输出文件路径。
        source_format: 源格式。
        target_format: 目标格式。

    返回:
        {"passed": bool, "issues": [str], "details": dict}
    """
    path = Path(output_path)
    issues: list[str] = []
    details = {}

    # 1. 文件存在
    if not path.exists():
        issues.append("输出文件不存在")
        return {"passed": False, "issues": issues, "details": details}

    # 2. 非空
    size = path.stat().st_size
    details["size_bytes"] = size
    if size == 0:
        issues.append("输出文件为空")

    # 3. 扩展名正确
    ext = path.suffix.lower().lstrip(".")
    expected_ext = _format_to_ext(target_format)
    details["extension"] = ext
    details["expected_extension"] = expected_ext
    if ext != expected_ext:
        issues.append(f"扩展名不匹配：实际 .{ext}，期望 .{expected_ext}")

    # 4. 内容完整性检查
    try:
        if target_format == "markdown":
            text = path.read_text(encoding="utf-8")
            details["char_count"] = len(text)
            if len(text.strip()) < 10:
                issues.append("Markdown 内容过少，可能转换失败")
            # 检查表格是否保留（如果源是 PDF/DOCX 且指令要求表格）
            if source_format in ("pdf", "docx"):
                table_lines = [l for l in text.splitlines() if l.strip().startswith("|")]
                details["table_line_count"] = len(table_lines)
        elif target_format == "docx":
            # DOCX 是二进制，检查大小合理
            if size < 1000:
                issues.append("DOCX 文件过小，可能内容丢失")
        elif target_format == "pdf":
            if size < 500:
                issues.append("PDF 文件过小，可能内容丢失")
        elif target_format == "xlsx":
            if size < 1000:
                issues.append("XLSX 文件过小，可能内容丢失")
            else:
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(path, read_only=True)
                    details["sheet_count"] = len(wb.sheetnames)
                    details["sheet_names"] = wb.sheetnames
                    wb.close()
                except Exception as e:
                    issues.append(f"XLSX 文件无法读取: {e}")
        elif target_format == "html":
            text = path.read_text(encoding="utf-8")
            details["char_count"] = len(text)
            if len(text.strip()) < 10:
                issues.append("HTML 内容过少，可能转换失败")
        elif target_format == "txt":
            text = path.read_text(encoding="utf-8")
            details["char_count"] = len(text)
            if len(text.strip()) < 10:
                issues.append("TXT 内容过少，可能转换失败")
    except Exception as e:
        issues.append(f"内容检查异常: {e}")

    return {
        "passed": len(issues) == 0,
        "issues": issues,
        "details": details,
    }


def _format_to_ext(fmt: str) -> str:
    """格式名转文件扩展名。"""
    return {
        "markdown": "md",
        "pdf": "pdf",
        "docx": "docx",
        "html": "html",
        "xlsx": "xlsx",
        "txt": "txt",
    }.get(fmt, fmt)
